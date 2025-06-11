import io
import pytest
import json
from starlette.datastructures import UploadFile
from FileConversion.converter import FileConversion

from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook

# -----------------------------
# File Generator Helpers
# -----------------------------

def make_pdf():
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    c.drawString(100, 750, "Hello PDF World")
    c.save()
    buffer.seek(0)
    return buffer

def make_docx():
    buffer = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Hello DOCX World")
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def make_xlsx():
    buffer = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Age"])
    ws.append(["Alice", 30])
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def make_csv():
    return io.BytesIO(b"name,age\nAlice,30\nBob,25")

def make_txt():
    return io.BytesIO(b"Just a plain text.")

def make_json():
    return io.BytesIO(json.dumps({"message": "Hello JSON World", "user": {"name": "Alice"}}).encode())

def make_xml():
    return io.BytesIO(b"<root><message>Hello XML World</message></root>")

def make_html():
    return io.BytesIO(b"<html><body><h1>Hello HTML</h1><p>World</p></body></html>")

VALID_TEST_FILES = [
    ("pdf", "application/pdf", make_pdf, "Hello PDF World"),
    ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", make_docx, "Hello DOCX World"),
    ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", make_xlsx, "Alice 30"),
    ("csv", "text/csv", make_csv, "Alice 30 Bob 25"),
    ("txt", "text/plain", make_txt, "Just a plain text."),
    ("json", "application/json", make_json, "Hello JSON World Alice"),
    ("xml", "application/xml", make_xml, "Hello XML World"),
    ("html", "text/html", make_html, "Hello HTML World"),
]

# -----------------------------
# Parametrized Tests
# -----------------------------

@pytest.mark.parametrize("ext, content_type, file_maker, expected", VALID_TEST_FILES)
def test_file_conversion_text_extraction(ext, content_type, file_maker, expected):
    file = UploadFile(filename=f"file.{ext}", file=file_maker())
    converter = FileConversion(file=file, extension=ext, content_type=content_type)

    text = converter.get_text()

    assert expected in text
